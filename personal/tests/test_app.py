import io
import re
import secrets
import sqlite3
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from personal.app import create_app, cents


class PersonalAppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({'TESTING': True, 'DATA_DIR': Path(self.temp.name), 'SECRET_KEY': 'test-session-only'})
        self.client = self.app.test_client()
        self.password = secrets.token_urlsafe(20)

    def tearDown(self):
        self.temp.cleanup()

    def token(self, path='/setup'):
        page = self.client.get(path).get_data(as_text=True)
        return re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)

    def post(self, path, data=None, token_path='/'):
        token = self.token(token_path)
        return self.client.post(path, data={'csrf_token': token, **({'username':'ana'} if path in ['/setup','/register','/login'] else {}), **(data or {})}, follow_redirects=True)

    def setup_owner(self):
        response = self.post('/setup', {'name': 'Ana', 'password': self.password, 'confirm': self.password}, '/setup')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Hola, Ana', response.get_data(as_text=True))

    def connection(self):
        return sqlite3.connect(self.app.config['DATABASE'])

    def new_client(self, **overrides):
        data = dict(name='Ana López', reference='CT-001', phone='5512345678', manager='Personal', balance='1,500.50')
        data.update(overrides)
        response = self.post('/clients/new', data)
        self.assertEqual(response.status_code, 200)
        with self.connection() as db:
            return db.execute('SELECT id FROM clients ORDER BY id DESC LIMIT 1').fetchone()[0]

    def import_csv(self, content, **overrides):
        token = self.token('/import')
        response = self.client.post('/import', data={'csrf_token': token, 'file': (io.BytesIO(content.encode()), 'cartera.csv')})
        self.assertEqual(response.status_code, 302)
        stage = response.headers['Location']
        mapping = dict(reference='Folio', name='Nombre', phone='Telefono', balance='Saldo', manager='Gestor', status='Estado', active_only='on')
        mapping.update(overrides)
        return self.post(stage, mapping, stage)

    def payment(self, client_id, amount, key=None):
        page = self.client.get(f'/clients/{client_id}').get_data(as_text=True)
        csrf = re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)
        key = key or re.search(r'name="payment_token" value="([^"]+)"', page).group(1)
        response = self.client.post(f'/clients/{client_id}/payments', data={'csrf_token': csrf, 'payment_token': key, 'amount': amount, 'note': 'Transferencia'}, follow_redirects=True)
        return response, key

    def test_initial_admin_login_logout_and_protected_routes(self):
        self.assertEqual(self.client.get('/clients').headers['Location'], '/setup')
        self.setup_owner()
        self.assertEqual(self.client.get('/setup').headers['Location'], '/login')
        self.post('/logout')
        self.assertEqual(self.client.get('/clients').headers['Location'], '/login')
        self.assertIn('no son correctos', self.post('/login', {'password': 'wrong'}, '/login').get_data(as_text=True))
        self.assertIn('Hola, Ana', self.post('/login', {'password': self.password}, '/login').get_data(as_text=True))
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM owner').fetchone()[0], 1)
            self.assertNotEqual(db.execute('SELECT password FROM owner').fetchone()[0], self.password)

    def test_csrf_blocks_state_changes(self):
        self.setup_owner()
        response = self.client.post('/clients/new', data={'name': 'No autorizado', 'balance': '100'})
        self.assertEqual(response.status_code, 400)
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM clients').fetchone()[0], 0)

    def test_payments_exact_cents_no_overpayment_and_no_duplicate_submission(self):
        self.setup_owner()
        client_id = self.new_client(balance='100.30')
        response, key = self.payment(client_id, '0.10')
        self.assertIn('Pago registrado', response.get_data(as_text=True))
        self.payment(client_id, '0.10', key)
        rejected, _ = self.payment(client_id, '101')
        self.assertIn('no puede superar', rejected.get_data(as_text=True))
        self.payment(client_id, '100.20')
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT balance_cents FROM clients').fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT COUNT(*), SUM(amount_cents) FROM events WHERE kind='payment'").fetchone(), (2, 10030))
        self.assertIn('Liquidado', self.client.get(f'/clients/{client_id}').get_data(as_text=True))

    def test_import_filters_and_reimport_preserves_payment(self):
        self.setup_owner()
        csv = 'Folio,Nombre,Telefono,Saldo,Gestor,Estado\nA,Ana,5512345678,1500,ANA,ACTIVO\nB,Beto,5598765432,2000,BETO,ACTIVO\nC,Caro,5512345670,0,ANA,PAGADO\nD,Dani,5512345671,100,ANA,LIQUIDADO\n'
        result = self.import_csv(csv, manager_filter='ANA')
        self.assertIn('1 clientes importados', result.get_data(as_text=True))
        with self.connection() as db:
            client_id = db.execute('SELECT id FROM clients').fetchone()[0]
        self.payment(client_id, '500')
        self.import_csv(csv, manager_filter='ANA')
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT COUNT(*),SUM(balance_cents) FROM clients').fetchone(), (1, 100000))
            self.assertEqual(db.execute("SELECT COUNT(*) FROM events WHERE kind='payment'").fetchone()[0], 1)

    def test_invalid_import_is_atomic(self):
        self.setup_owner()
        result = self.import_csv('Folio,Nombre,Telefono,Saldo,Gestor,Estado\nA,Ana,5512345678,100,ANA,ACTIVO\nB,Beto,5512345679,no-es-un-monto,ANA,ACTIVO\n')
        self.assertIn('No se importó ningún cliente', result.get_data(as_text=True))
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM clients').fetchone()[0], 0)

    def test_xlsx_import_named_sheet_and_manual_mapping(self):
        self.setup_owner()
        file = io.BytesIO()
        with pd.ExcelWriter(file, engine='openpyxl') as writer:
            pd.DataFrame({'Ayuda': ['No importar']}).to_excel(writer, sheet_name='Instrucciones', index=False)
            pd.DataFrame({'Persona': ['Ana'], 'Pendiente': ['1.234,56']}).to_excel(writer, sheet_name='Cartera', index=False)
        token = self.token('/import')
        response = self.client.post('/import', data={'csrf_token': token, 'sheet': 'Cartera', 'file': (io.BytesIO(file.getvalue()), 'cartera.xlsx')})
        stage = response.headers['Location']
        result = self.post(stage, {'name': 'Persona', 'balance': 'Pendiente'}, stage)
        self.assertEqual(result.status_code, 200)
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT name,balance_cents FROM clients').fetchone(), ('Ana', 123456))

    def test_filtered_exports_match_clients_and_do_not_create_formulas(self):
        self.setup_owner()
        self.new_client(name='=HYPERLINK("https://example.com")', manager='ANA')
        self.new_client(name='Beto', reference='CT-002', manager='BETO')
        response = self.client.get('/export?manager=ANA')
        workbook = load_workbook(io.BytesIO(response.data))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertTrue(sheet['B2'].value.startswith("'="))
        csv_response = self.client.get('/export?manager=BETO&format=csv')
        self.assertIn('Beto', csv_response.data.decode('utf-8-sig'))
        self.assertNotIn('HYPERLINK', csv_response.data.decode('utf-8-sig'))
        self.assertEqual(self.client.get('/clients?q=Beto').get_data(as_text=True).count('CT-002'), 1)

    def test_duplicate_reference_does_not_overwrite_customer(self):
        self.setup_owner()
        self.new_client()
        self.new_client(name='Otro nombre', balance='1')
        with self.connection() as db:
            self.assertEqual(db.execute('SELECT COUNT(*),name,balance_cents FROM clients').fetchone(), (1, 'Ana López', 150050))

    def test_saved_state_survives_new_app_instance(self):
        self.setup_owner()
        client_id = self.new_client()
        self.payment(client_id, '500')
        restored = create_app({'TESTING': True, 'DATA_DIR': Path(self.temp.name), 'SECRET_KEY': 'test-session-only'})
        with sqlite3.connect(restored.config['DATABASE']) as db:
            self.assertEqual(db.execute('SELECT balance_cents FROM clients').fetchone()[0], 100050)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM owner').fetchone()[0], 1)

    def test_notes_escape_html_and_history_records_changes(self):
        self.setup_owner()
        client_id = self.new_client()
        response = self.post(f'/clients/{client_id}', {'name': '<script>alert(1)</script>', 'phone': '5511111111', 'manager': 'Personal', 'notes': 'Llamar el viernes'})
        html = response.get_data(as_text=True)
        self.assertNotIn('<script>alert(1)</script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertIn('Datos guardados', html)
        self.assertIn('Datos de contacto y notas actualizados', self.client.get('/history').get_data(as_text=True))

    def test_login_rate_limit(self):
        self.setup_owner()
        self.post('/logout')
        for _ in range(10):
            self.post('/login', {'password': 'wrong'}, '/login')
        response = self.post('/login', {'password': self.password}, '/login')
        self.assertEqual(response.status_code, 429)

    def test_amount_formats_and_reject_negative_or_nonfinite(self):
        for raw, expected in [('1,234.56', 123456), ('1.234,56', 123456), ('$ 1,000', 100000), ('1500,50', 150050), ('0.10', 10)]:
            self.assertEqual(cents(raw), expected)
        for raw in ['-10', 'NaN', 'Infinity', '', 'incorrecto', '1,,2', '1,2.30', '1.234']:
            with self.assertRaises(ValueError):
                cents(raw)


if __name__ == '__main__':
    unittest.main()
